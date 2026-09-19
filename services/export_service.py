"""基于一致快照的 Excel 报表。所有外部文本都以字符串单元格保存。"""

from __future__ import annotations

import asyncio
import os
import uuid
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from ..utils import visual_style as V
from ..utils.time_utils import parse_iso, to_beijing
from .report_service import ReportService


@dataclass(slots=True)
class ExportBundle:
    path: Path | None
    filename: str
    empty: bool = False


class ExportService:
    def __init__(
        self,
        *,
        user_repo,
        submission_repo,
        admin_repo,
        pause_repo,
        export_dir,
        weekly_limit,
        timezone,
    ):
        self.report = ReportService(user_repo.db, weekly_limit, timezone)
        self.export_dir, self.timezone = Path(export_dir), timezone
        self._semaphore = asyncio.Semaphore(1)

    async def export(self, user_id=None):
        snapshot = await self.report.snapshot(user_id)
        if not snapshot.users:
            return ExportBundle(None, "", True)
        async with self._semaphore:
            return await asyncio.to_thread(self.render, snapshot)

    def _date(self, value):
        moment = parse_iso(value) if isinstance(value, str) else value
        return to_beijing(moment, self.timezone).replace(tzinfo=None) if moment else None

    def render(self, snapshot):
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.chart import DoughnutChart, Reference
        from openpyxl.chart.series import DataPoint
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        overview = workbook.active
        overview.title = "本周总览"
        font_name = "Microsoft YaHei"
        navy, teal, blue = V.NAVY[1:], V.TEAL[1:], V.BLUE[1:]
        fills = {
            "已完成": "DFF3EF",
            "进行中": "E4EFFB",
            "未开始": "EDF1F5",
            "暂停周": "FFF2D9",
            "有效计次": "DFF3EF",
            "额外材料": "E4EFFB",
            "审核服务异常": "FCE7E7",
            "处理异常": "FCE7E7",
        }

        def write(sheet, row, values):
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row, column)
                if isinstance(value, str):
                    cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
                    cell.data_type = "s"
                    cell.number_format = "@"
                    cell.quotePrefix = True
                else:
                    cell.value = value
                    cell.number_format = (
                        "yyyy-mm-dd hh:mm" if isinstance(value, datetime) else "#,##0"
                    )

        def style_table(sheet, headers, rows, widths, status_columns=(), freeze="D2"):
            write(sheet, 1, headers)
            for index, values in enumerate(rows, 2):
                write(sheet, index, values)
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = freeze
            sheet.auto_filter.ref = sheet.dimensions
            sheet.row_dimensions[1].height = 42
            for col, width in enumerate(widths, 1):
                sheet.column_dimensions[get_column_letter(col)].width = width
            for row in sheet:
                for cell in row:
                    cell.font = Font(name=font_name, size=11, color=navy)
                    cell.alignment = Alignment(
                        vertical="center",
                        horizontal="left"
                        if cell.data_type == "s" and not str(cell.value or "").isdecimal()
                        else "center",
                        indent=1,
                        wrap_text=True,
                    )
                    if cell.row == 1:
                        cell.fill = PatternFill("solid", fgColor=navy)
                        cell.font = Font(name=font_name, size=11, bold=True, color="FFFFFF")
                        cell.alignment = Alignment(
                            horizontal="center", vertical="center", wrap_text=True
                        )
                    elif cell.row % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor="F2F7FB")
                    if cell.row > 1 and cell.column in status_columns and cell.value in fills:
                        cell.fill = PatternFill("solid", fgColor=fills[cell.value])
                if row[0].row > 1:
                    line_count = max(
                        1,
                        max(
                            (
                                sum(2 if ord(c) > 255 else 1 for c in str(cell.value or ""))
                                // max(1, int(widths[cell.column - 1]) - 2)
                                + str(cell.value or "").count("\n")
                                + 1
                            )
                            for cell in row
                        ),
                    )
                    sheet.row_dimensions[row[0].row].height = min(409, max(28, line_count * 17))
            for column in status_columns:
                letter = get_column_letter(column)
                for label, color in fills.items():
                    if sheet.max_row > 1:
                        sheet.conditional_formatting.add(
                            f"{letter}2:{letter}{sheet.max_row}",
                            CellIsRule(
                                operator="equal",
                                formula=['"' + label + '"'],
                                fill=PatternFill("solid", fgColor=color),
                            ),
                        )
            sheet.print_title_rows = "1:1"
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
            sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
            sheet.print_options.horizontalCentered = True
            sheet.oddFooter.center.text = "第 &P 页 / 共 &N 页"
            sheet.sheet_properties.outlinePr.summaryRight = False

        def completion(row):
            if row.user.status != "active":
                return "不参与考核"
            if snapshot.paused:
                return "暂停周"
            return (
                "已完成"
                if row.current >= snapshot.target
                else "进行中"
                if row.current
                else "未开始"
            )

        summary = workbook.create_sheet("用户汇总")
        headers = [
            "QQ号",
            "学号",
            "姓名",
            "本周状态",
            "本周有效计次",
            "本周目标",
            "本周自动计次",
            "本周人工计次",
            "本周人工额外",
            "本周扣减",
            "累计周有效计次",
            "有效材料数",
            "其中额外材料",
            "累计人工计次",
            "累计人工额外",
            "累计扣减",
            "累计人工调整净额",
            "内容或文件未通过",
            "系统异常",
            "完成周数",
            "豁免周数",
            "绑定时间",
            "最后有效提交",
            "用户状态",
        ]
        rows = [
            [
                r.user.qq_id,
                r.user.student_id,
                r.user.name,
                completion(r),
                r.current,
                0 if snapshot.paused else snapshot.target,
                r.automatic,
                r.manual_counted,
                r.manual_extra,
                r.deductions,
                r.cumulative_counted,
                r.materials,
                r.extra_materials,
                r.all_manual_counted,
                r.all_manual_extra,
                r.all_deductions,
                r.adjustments_net,
                r.rejected,
                r.errors,
                r.completed_weeks,
                r.exempt_weeks,
                self._date(r.user.bound_at),
                self._date(r.user.last_submission_at),
                "在队" if r.user.status == "active" else "停用",
            ]
            for r in snapshot.users
        ]
        style_table(summary, headers, rows, [17, 20, 18, 16] + [15] * 17 + [22, 22, 14], (4,))
        summary.sheet_properties.tabColor = teal
        for start, end in (("G", "J"), ("N", "S"), ("V", "X")):
            summary.column_dimensions.group(start, end, hidden=True)

        detail = workbook.create_sheet("打卡明细")
        headers = [
            "QQ号",
            "学号",
            "姓名",
            "提交时间",
            "周起始日期",
            "状态",
            "是否计次",
            "原文件名",
            "AI 简评",
            "相似度",
            "AI 置信度",
            "审核结论",
            "文件大小（MiB）",
            "submission_id",
            "群或会话 ID",
            "原消息 ID",
            "引用消息 ID",
            "NAS 保存路径",
            "SHA256",
            "文本哈希",
            "重复检测类型",
            "AI Provider ID",
            "AI 耗时（ms）",
            "错误信息",
        ]
        rows = [
            [
                s.qq_id_snapshot,
                s.student_id_snapshot,
                s.name_snapshot,
                self._date(s.submitted_at),
                datetime.fromisoformat(s.week_key),
                V.status_label(s.status),
                "是" if s.counted else "否",
                s.original_filename,
                s.ai_feedback or "",
                s.similarity_score,
                s.ai_confidence,
                {"pass": "通过", "fail": "未通过"}.get(s.ai_decision, "未完成"),
                s.file_size / 1048576,
                s.id,
                s.qq_group_id or "",
                s.qq_message_id,
                s.quoted_message_id,
                s.stored_path,
                s.sha256,
                s.normalized_text_sha256 or "",
                s.duplicate_type,
                s.ai_provider_id or "",
                s.ai_latency_ms,
                s.error_message or "",
            ]
            for s in snapshot.submissions
        ]
        style_table(
            detail,
            headers,
            rows,
            [
                17,
                20,
                18,
                22,
                17,
                18,
                12,
                30,
                52,
                13,
                13,
                14,
                18,
                36,
                20,
                22,
                22,
                65,
                65,
                65,
                24,
                28,
                18,
                52,
            ],
            (6,),
        )
        detail.column_dimensions.group("N", "W", hidden=True)
        for row in detail.iter_rows(min_row=2):
            row[4].number_format = "yyyy-mm-dd"
            row[9].number_format = row[10].number_format = "0.0%"
            row[12].number_format = "0.00"

        user_by_id = {r.user.id: r.user for r in snapshot.users}
        adjustments = workbook.create_sheet("人工调整")
        rows = []
        for a in snapshot.adjustments:
            user = user_by_id[a.user_id]
            rows.append(
                [
                    self._date(a.created_at),
                    datetime.fromisoformat(a.week_key),
                    user.qq_id,
                    user.student_id,
                    user.name,
                    a.delta,
                    "扣减" if a.delta < 0 else "计入本周" if a.counted else "额外记录",
                    a.admin_qq,
                    a.reason or "",
                ]
            )
        style_table(
            adjustments,
            [
                "时间",
                "周起始日期",
                "目标 QQ",
                "学号",
                "姓名",
                "调整值",
                "计次方式",
                "管理员 QQ",
                "原因",
            ],
            rows,
            [22, 17, 17, 20, 18, 12, 18, 18, 45],
            freeze="F2",
        )
        for row in adjustments.iter_rows(min_row=2):
            row[1].number_format = "yyyy-mm-dd"
            row[5].number_format = "+0;-0;0"

        pauses = workbook.create_sheet("暂停记录")
        rows = [
            [
                {"day": "当天", "week": "本周"}.get(p.scope, p.scope),
                self._date(p.start_at),
                self._date(p.scheduled_end_at),
                self._date(p.resumed_at),
                p.reason,
                p.created_by_qq,
                p.exemption_week_key or "",
                "已取消豁免"
                if p.resumed_at and p.exemption_week_key
                else "豁免周"
                if p.exemption_week_key
                else "仅暂停接收",
            ]
            for p in snapshot.pauses
        ]
        style_table(
            pauses,
            [
                "范围",
                "开始时间",
                "计划结束（不含）",
                "实际恢复时间",
                "原因",
                "管理员 QQ",
                "关联考核周",
                "豁免状态",
            ],
            rows,
            [12, 22, 22, 22, 48, 18, 18, 18],
            freeze="B2",
        )

        overview.sheet_view.showGridLines = False
        overview.sheet_properties.tabColor = navy
        for column in "ABCDEFGH":
            overview.column_dimensions[column].width = 16
        overview.column_dimensions["A"].width = 20
        overview.column_dimensions["B"].width = 22
        overview.column_dimensions["C"].width = 18
        for column in "DEF":
            overview.column_dimensions[column].width = 14
        for row in range(1, 22):
            overview.row_dimensions[row].height = 26
        title = "个人学习记录" if snapshot.personal else "直属队学习打卡"
        write(overview, 2, [title])
        overview["A2"].font = Font(name=font_name, size=20, color=navy, bold=True)
        write(
            overview,
            3,
            [
                f"考核周 {snapshot.week_key} 起    每人目标 {snapshot.target} 次    导出时间 {self._date(snapshot.generated_at):%Y-%m-%d %H:%M}"
            ],
        )
        overview["A3"].font = Font(name=font_name, size=10, color=V.MUTED[1:])
        completed, ongoing, zero = snapshot.counts
        write(
            overview,
            5,
            [
                "参与人数",
                snapshot.total,
                "已完成人数",
                completed if not snapshot.paused else "—",
                "完成率",
                completed / snapshot.total if snapshot.total and not snapshot.paused else "—",
            ],
        )
        overview["F5"].number_format = "0%"
        for cell in overview[5][:6]:
            cell.font = Font(
                name=font_name, size=14, color=teal if cell.column % 2 == 0 else navy, bold=True
            )
            cell.fill = PatternFill("solid", fgColor="E6F4F4")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        if snapshot.paused:
            write(overview, 7, ["本周暂停考核，已有记录保留；本周不按缺卡处理。"])
            pause_text = "暂停原因：" + snapshot.reason
            write(overview, 9, [pause_text])
            overview.merge_cells("A9:H11")
            overview["A9"].alignment = Alignment(wrap_text=True, vertical="top")
            overview.row_dimensions[9].height = max(26, ((len(pause_text) * 2 // 130) + 1) * 17)
        elif snapshot.total:
            write(overview, 7, ["本周分布", "人数"])
            for row, values in enumerate(
                zip(("已完成", "进行中", "未开始"), snapshot.counts, strict=True), 8
            ):
                write(overview, row, values)
            chart = DoughnutChart()
            chart.title, chart.holeSize = "本周完成情况", 72
            chart.add_data(
                Reference(overview, min_col=2, min_row=7, max_row=10), titles_from_data=True
            )
            chart.set_categories(Reference(overview, min_col=1, min_row=8, max_row=10))
            chart.height, chart.width = 7.2, 13.5
            chart.legend.position = "r"
            for idx, color in enumerate((teal, blue, V.ZERO[1:])):
                point = DataPoint(idx=idx)
                point.graphicalProperties.solidFill = color
                chart.series[0].data_points.append(point)
            overview.add_chart(chart, "D7")
        else:
            write(overview, 7, ["该用户已停用，不参与当前考核。"])
        for row, label in (
            (13, "累计周有效计次：每周有效净次数按该周目标封顶，再逐周相加。"),
            (14, "有效材料数：审核通过的全部材料，包含同日及超周目标的额外材料。"),
            (15, "人工计次、人工额外和扣减单列；额外材料不抵扣其他周缺卡。"),
            (16, "统计按在队用户计算；豁免周单列，不作为正常完成周。"),
        ):
            write(overview, row, [label])
            overview.cell(row, 1).font = Font(name=font_name, size=10, color=V.MUTED[1:])
        # 说明置于图表下方，避免与图形重叠。
        for source, destination in ((13, 19), (14, 20), (15, 21), (16, 22)):
            overview.cell(destination, 1, overview.cell(source, 1).value)
            overview.cell(destination, 1).font = copy(overview.cell(source, 1).font)
            overview.cell(source, 1).value = None
        write(overview, 24, ["个人进度" if snapshot.personal else "本周待完成名单"])
        write(overview, 25, ["学号", "姓名", "状态", "已计次", "本周目标", "还需完成"])
        candidates = (
            snapshot.users
            if snapshot.personal
            else [
                r
                for r in snapshot.population
                if r.current < snapshot.target and not snapshot.paused
            ]
        )
        for row, item in enumerate(
            sorted(candidates, key=lambda r: (r.current, r.user.student_id)), 26
        ):
            required = 0 if snapshot.paused or item.user.status != "active" else snapshot.target
            write(
                overview,
                row,
                [
                    item.user.student_id,
                    item.user.name,
                    completion(item),
                    item.current,
                    required,
                    max(0, required - item.current),
                ],
            )
            overview.row_dimensions[row].height = max(
                30, ((len(item.user.name) * 2 // 20) + 1) * 17
            )
        if not candidates:
            write(
                overview,
                26,
                ["本周豁免，无待完成名单。" if snapshot.paused else "当前没有待完成的同学。"],
            )
        for row in overview.iter_rows(min_row=24):
            for cell in row:
                cell.font = Font(name=font_name, size=11, color=navy, bold=cell.row in {24, 25})
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                    indent=1,
                    horizontal="left" if cell.data_type == "s" else "center",
                )
                if cell.row == 25:
                    cell.fill = PatternFill("solid", fgColor="E4EFFB")
                elif cell.row > 25 and cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F2F7FB")
        # 总览保持自由滚动；长名单可在用户汇总筛选，不冻结整屏标题区。
        overview.freeze_panes = None
        overview.print_options.horizontalCentered = True
        overview.page_setup.orientation = "landscape"
        overview.page_setup.paperSize = overview.PAPERSIZE_A4
        overview.sheet_properties.pageSetUpPr.fitToPage = True
        overview.page_setup.fitToWidth, overview.page_setup.fitToHeight = 1, 0
        overview.print_area = f"A1:H{max(26, overview.max_row)}"
        overview.print_title_rows = "25:25"
        overview["A2"].border = Border(bottom=Side(style="thin", color=teal))
        self.export_dir.mkdir(parents=True, exist_ok=True)
        suffix = uuid.uuid4().hex[:12]
        filename = f"直属队打卡_{snapshot.week_key}_{'个人' if snapshot.personal else '全员'}_{suffix}.xlsx"
        path = self.export_dir / filename
        partial = path.with_suffix(".part")
        try:
            workbook.save(partial)
            os.replace(partial, path)
        finally:
            workbook.close()
            partial.unlink(missing_ok=True)
        return ExportBundle(path, filename)
