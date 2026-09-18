"""Word（.docx）校验与内容提取。

安全要求：
- 只接受有效的 Office Open XML（ZIP）文件；
- 限制解压后总大小与压缩比，防止 ZIP Bomb；
- 不执行文档中的宏、脚本或嵌入程序。
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from ..utils.text_utils import normalize_text, select_relevant_text


class DocxValidationError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


class DocxParseError(Exception):
    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.message = message


@dataclass(slots=True)
class DocxContent:
    text: str = ""
    char_count: int = 0
    table_count: int = 0
    image_count: int = 0
    link_count: int = 0
    headings: list[str] = field(default_factory=list)
    created: str = ""
    modified: str = ""


def validate_docx(path: Path, max_uncompressed_bytes: int) -> None:
    """校验 docx 结构，失败时抛出 DocxValidationError。"""

    path = Path(path)
    if not path.is_file():
        raise DocxValidationError("E_DOCX_PARSE_FAILED", "文件不存在")
    if not zipfile.is_zipfile(path):
        raise DocxValidationError("E_DOCX_PARSE_FAILED", "不是有效的 Office Open XML 文件")

    try:
        with zipfile.ZipFile(path) as zf:
            names = set(zf.namelist())
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocxValidationError("E_DOCX_PARSE_FAILED", "缺少 Word 文档结构")
            total = 0
            for info in zf.infolist():
                total += info.file_size
                if total > max_uncompressed_bytes:
                    raise DocxValidationError("E_FILE_TOO_LARGE", "解压后内容过大，疑似异常文件")
                compressed = max(info.compress_size, 1)
                if info.file_size > 5 * 1024 * 1024 and info.file_size / compressed > 200:
                    raise DocxValidationError("E_FILE_TOO_LARGE", "压缩比异常，疑似 ZIP Bomb")
            if zf.testzip() is not None:
                raise DocxValidationError("E_DOCX_PARSE_FAILED", "压缩包内容损坏")
    except zipfile.BadZipFile as exc:
        raise DocxValidationError("E_DOCX_PARSE_FAILED", "压缩包无法解析") from exc


def extract_content(path: Path) -> DocxContent:
    """提取正文、表格文字与结构证据。"""

    path = Path(path)
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - 依赖缺失时的保护
        raise DocxParseError("服务端缺少 python-docx 依赖") from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise DocxParseError("无法打开 Word 文档") from exc

    segments: list[str] = []
    headings: list[str] = []
    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        segments.append(text)
        style_name = ""
        try:
            style_name = (paragraph.style.name or "") if paragraph.style else ""
        except Exception:
            style_name = ""
        if style_name.lower().startswith("heading") or style_name.startswith("标题"):
            headings.append(text[:100])

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                segments.append(" | ".join(cells))

    text = normalize_text("\n".join(segments))

    image_count = 0
    link_count = 0
    try:
        rels = document.part.rels
        for rel in rels.values():
            reltype = getattr(rel, "reltype", "") or ""
            if "image" in reltype:
                image_count += 1
            elif "hyperlink" in reltype:
                link_count += 1
    except Exception:
        pass

    created = ""
    modified = ""
    try:
        props = document.core_properties
        created = props.created.isoformat() if props.created else ""
        modified = props.modified.isoformat() if props.modified else ""
    except Exception:
        pass

    return DocxContent(
        text=text,
        char_count=len(text),
        table_count=len(document.tables),
        image_count=image_count,
        link_count=link_count,
        headings=headings[:20],
        created=created,
        modified=modified,
    )


def text_for_ai(content: DocxContent, max_chars: int) -> str:
    return select_relevant_text(content.text, max_chars)


def evidence_text(content: DocxContent) -> str:
    return (
        f"字符数={content.char_count}，表格数={content.table_count}，"
        f"图片数={content.image_count}，链接数={content.link_count}"
    )
