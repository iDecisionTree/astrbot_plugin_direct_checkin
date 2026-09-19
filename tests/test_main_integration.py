"""AstrBot 接口替身覆盖完整命令入口；不代表真实 QQ 平台验收。"""

import asyncio
import importlib
import logging
import sys
import types
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest


class File:
    def __init__(self, file="", name="", url=""):
        self.file, self.name, self.url = file, name, url

    async def get_file(self):
        return self.file


class Reply:
    def __init__(self, chain):
        self.chain, self.id = chain, "quoted"


class Event:
    def __init__(self, text, message_id, sender="9", bot="1", chain=None):
        self.message_str = text
        self.message_obj = SimpleNamespace(message_id=message_id, self_id=bot, message=chain or [])
        self.unified_msg_origin = "aiocqhttp:GroupMessage:3"
        self.sender = sender
        self.stopped = False

    def get_sender_id(self):
        return self.sender

    def get_group_id(self):
        return "3"

    def stop_event(self):
        self.stopped = True

    def plain_result(self, text):
        return SimpleNamespace(chain=[SimpleNamespace(text=text)])

    def image_result(self, path):
        return SimpleNamespace(chain=[SimpleNamespace(text="", file=path)])

    def chain_result(self, chain):
        return SimpleNamespace(chain=chain)


@pytest.fixture
def plugin(tmp_path, monkeypatch):
    modules = {
        name: types.ModuleType(name)
        for name in (
            "astrbot",
            "astrbot.api",
            "astrbot.api.event",
            "astrbot.api.star",
            "astrbot.api.message_components",
        )
    }
    logger = logging.getLogger("test-plugin")
    modules["astrbot.api"].logger = logger

    class Star:
        def __init__(self, context, config):
            self.context, self.logger = context, logger

    def decorator(*args, **kwargs):
        return lambda fn: fn

    modules["astrbot.api.event"].AstrMessageEvent = Event
    modules["astrbot.api.event"].filter = SimpleNamespace(
        command=decorator,
        platform_adapter_type=decorator,
        PlatformAdapterType=SimpleNamespace(AIOCQHTTP="qq"),
    )
    modules["astrbot.api.star"].Star = Star
    modules["astrbot.api.star"].Context = object
    modules["astrbot.api.star"].register = decorator
    modules["astrbot.api.message_components"].File = File
    modules["astrbot.api.message_components"].Reply = Reply
    for name, module in modules.items():
        monkeypatch.setitem(sys.modules, name, module)
    name = "astrbot_plugin_direct_checkin.main"
    monkeypatch.delitem(sys.modules, name, raising=False)
    main = importlib.import_module(name)

    def paths(self):
        self.data_dir = tmp_path / "data"
        self.data_dir.mkdir()
        self.export_dir = self.data_dir / "exports"
        self.image_dir = self.data_dir / "images"

    monkeypatch.setattr(main.DirectCheckinPlugin, "_setup_paths", paths)

    class Context:
        async def get_current_chat_provider_id(self, umo):
            return "test"

        async def llm_generate(self, **kwargs):
            return SimpleNamespace(
                completion_text='{"decision":"pass","brief_feedback":"实现清晰并有测试验证"}'
            )

    nas = tmp_path / "nas"
    nas.mkdir()
    return main.DirectCheckinPlugin(
        Context(), {"nas_base_dir": str(nas), "bootstrap_admin_qq": "9"}
    ), main


async def send(plugin, event):
    items = [result async for result in plugin.direct_checkin(event)]
    assert event.stopped
    return "\n".join(getattr(c, "text", "") for item in items for c in item.chain)


def test_entry_binding_admin_idempotency_and_permission(plugin):
    app, _ = plugin

    async def run():
        await app.initialize()
        first = await send(app, Event("/d 20260001 张三", "bind", sender="10"))
        assert first == await send(app, Event("/d 20260001 张三", "bind", sender="10"))
        one = await send(app, Event("/d add sid:20260001", "add"))
        assert "1/2" in one
        assert one == await send(app, Event("/d add sid:20260001", "add"))
        two = await send(app, Event("/d add sid:20260001", "add", bot="2"))
        assert "2/2" in two
        assert len(await app.admin_repo.list_adjustments()) == 2
        assert "只有" in await send(app, Event("/d remove sid:20260001", "denied", sender="10"))
        assert "唯一标识" in await send(app, Event("/d remove sid:20260001", "", sender="9"))
        await app.terminate()

    asyncio.run(run())


def test_entry_checkin_uses_one_received_time_and_replays(plugin, tmp_path, monkeypatch):
    app, main = plugin
    import docx

    document = docx.Document()
    document.add_paragraph("本周实现文件上传接口，新增超限检查与测试，解决空文件报错。")
    source = tmp_path / "borrowed.docx"
    document.save(source)

    async def run():
        await app.initialize()
        await send(app, Event("/d 20260001 张三", "bind", sender="10"))
        # 处理跨越北京周日午夜，归属仍来自本次接收时间。
        before = datetime.fromisoformat("2026-09-20T15:59:59+00:00")
        after = datetime.fromisoformat("2026-09-20T16:01:00+00:00")
        monkeypatch.setattr(main, "now_utc", lambda: before)
        monkeypatch.setattr(
            "astrbot_plugin_direct_checkin.services.checkin_service.now_utc", lambda: after
        )
        event = Event("/d", "checkin", sender="10", chain=[Reply([File(str(source), "学习.docx")])])
        first = await send(app, event)
        assert "1/2" in first
        assert await send(app, event) == first
        records = await app.submission_repo.list_submissions()
        assert len(records) == 1
        assert records[0].week_key == "2026-09-14"
        assert records[0].beijing_date == "2026-09-20"
        assert source.exists() and Path(records[0].stored_path).exists()
        await app.terminate()

    asyncio.run(run())


def test_entry_reset_duplicate_confirmation_and_receipts_survive(plugin):
    app, _ = plugin

    async def run():
        await app.initialize()
        await send(app, Event("/d 20260001 张三", "bind", sender="10"))
        first = await send(app, Event("/d reset confirm", "r1"))
        assert first == await send(app, Event("/d reset confirm", "r1"))
        assert len(await app.user_repo.list_users()) == 1
        done = await send(app, Event("/d reset confirm", "r2"))
        assert "重置完成" in done and await app.user_repo.list_users() == []
        assert done == await send(app, Event("/d reset confirm", "r2"))
        assert await app.admin_repo.is_admin("9")
        await app.terminate()

    asyncio.run(run())


def test_committed_checkin_survives_reply_failure(plugin, tmp_path, monkeypatch):
    app, _ = plugin
    import docx

    document = docx.Document()
    document.add_paragraph("实现了输入校验，使用边界测试验证错误响应，修复空值缺陷。")
    source = tmp_path / "late.docx"
    document.save(source)

    def broken_reply(*args):
        raise RuntimeError("reply formatting")

    async def run():
        await app.initialize()
        await send(app, Event("/d 20260001 张三", "bind", sender="10"))
        monkeypatch.setattr(
            "astrbot_plugin_direct_checkin.utils.response_templates.checkin_success", broken_reply
        )
        response = await send(
            app, Event("/d", "late", sender="10", chain=[Reply([File(str(source), "学习.docx")])])
        )
        assert "已完成计次" in response
        records = await app.submission_repo.list_submissions()
        assert records[0].status == "VALID_COUNTED"
        assert (await app.user_repo.get_by_qq("10")).last_submission_at is not None
        await app.terminate()

    asyncio.run(run())


def test_failed_pending_insert_removes_only_own_archive(plugin, tmp_path, monkeypatch):
    app, _ = plugin
    from unittest.mock import AsyncMock

    import docx

    document = docx.Document()
    document.add_paragraph("实现输入校验并完成验证。")
    source = tmp_path / "source.docx"
    document.save(source)

    async def run():
        await app.initialize()
        await send(app, Event("/d 20260001 张三", "bind", sender="10"))
        monkeypatch.setattr(
            app.submission_repo, "create_pending", AsyncMock(side_effect=RuntimeError("db"))
        )
        response = await send(
            app, Event("/d", "failed", sender="10", chain=[Reply([File(str(source), "学习.docx")])])
        )
        assert "技术问题" in response and source.is_file()
        assert list(app.file_service.nas_base_dir.rglob("*.docx")) == []
        assert await app.submission_repo.list_submissions() == []
        await app.terminate()

    asyncio.run(run())


def test_entry_daily_order_matches_prompt_and_extra_reply(plugin, tmp_path, monkeypatch):
    app, main = plugin
    import json
    from unittest.mock import AsyncMock

    import docx

    source = tmp_path / "daily.docx"

    def document(text):
        file = docx.Document()
        file.add_paragraph(text)
        file.save(source)

    async def run():
        await app.initialize()
        await send(app, Event("/d 20260001 张三", "bind1", sender="10"))
        await send(app, Event("/d 20260002 李四", "bind2", sender="11"))
        before = datetime.fromisoformat("2026-09-20T15:59:59+00:00")
        monkeypatch.setattr(main, "now_utc", lambda: before)
        monkeypatch.setattr(
            "astrbot_plugin_direct_checkin.services.checkin_service.now_utc",
            lambda: datetime.fromisoformat("2026-09-20T16:05:00+00:00"),
        )
        model = AsyncMock(wraps=app.context.llm_generate)
        monkeypatch.setattr(app.context, "llm_generate", model)
        document("实现上传校验与边界测试，修复了空文件异常。")
        first_event = Event(
            "/d", "first", sender="10", chain=[Reply([File(str(source), "记录.docx")])]
        )
        first = await send(app, first_event)
        assert "第 1 位" in first and "1/2" in first
        assert await send(app, first_event) == first
        assert model.await_count == 1
        document("实现检索索引与性能测试，记录不同数据量的耗时和改进。")
        second = await send(
            app, Event("/d", "second", sender="11", chain=[Reply([File(str(source), "记录.docx")])])
        )
        assert "第 2 位" in second
        document("新增缓存失效策略，并完成并发测试，解决旧数据未刷新的问题。")
        extra = await send(
            app, Event("/d", "extra", sender="10", chain=[Reply([File(str(source), "记录.docx")])])
        )
        assert "第 1 位" in extra and "额外材料" in extra
        contexts = [
            json.loads(call.kwargs["prompt"])["checkin_context"] for call in model.call_args_list
        ]
        assert [c["daily_position"] for c in contexts] == [1, 2, 1]
        assert {c["date"] for c in contexts} == {"2026-09-20"}
        await app.terminate()

    asyncio.run(run())
