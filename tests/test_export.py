import asyncio
from datetime import datetime
from pathlib import Path

from astrbot_plugin_direct_checkin.models.enums import SubmissionStatus
from astrbot_plugin_direct_checkin.repositories.admin_repo import AdminRepository
from astrbot_plugin_direct_checkin.repositories.db import Database
from astrbot_plugin_direct_checkin.repositories.pause_repo import PauseRepository
from astrbot_plugin_direct_checkin.repositories.submission_repo import SubmissionRepository
from astrbot_plugin_direct_checkin.repositories.user_repo import UserRepository
from astrbot_plugin_direct_checkin.services.export_service import ExportService
from astrbot_plugin_direct_checkin.utils.time_utils import week_key_now
from openpyxl import load_workbook

NOW = "2026-01-05T02:00:00+00:00"


def test_export_workbook_structure(tmp_path: Path):
    async def scenario() -> None:
        db = Database(tmp_path / "checkin.db")
        await db.initialize()
        users = UserRepository(db)
        submissions = SubmissionRepository(db)
        admins = AdminRepository(db)
        pauses = PauseRepository(db)

        week_key = week_key_now("Asia/Shanghai")
        user = await users.create("1001", "2026123456", "张三", None, NOW)
        await submissions.create_pending(
            submission_id="s1",
            user_id=user.id,
            qq_id=user.qq_id,
            student_id=user.student_id,
            name=user.name,
            submitted_at=NOW,
            week_key=week_key,
            beijing_date="2026-01-05",
            qq_group_id="999",
            qq_message_id="msg-1",
            quoted_message_id="q1",
            original_filename="week.docx",
            stored_path="/mnt/nas/direct_checkin/2026123456/week.docx",
            file_size=1234,
            sha256="deadbeef",
            now=NOW,
        )
        await submissions.update(
            "s1",
            status=SubmissionStatus.VALID_COUNTED.value,
            counted=1,
            counted_slot=1,
            ai_decision="pass",
            updated_at=NOW,
        )
        await admins.allocate_positive_adjustment(
            user.id, week_key, "2026-01-06", 2, "2747344390", NOW
        )

        service = ExportService(
            user_repo=users,
            submission_repo=submissions,
            admin_repo=admins,
            pause_repo=pauses,
            export_dir=tmp_path / "exports",
            weekly_limit=2,
            timezone="Asia/Shanghai",
        )
        bundle = await service.export()
        assert bundle.empty is False
        assert bundle.path is not None and bundle.path.exists()

        workbook = load_workbook(bundle.path)
        assert workbook.sheetnames == ["本周总览", "用户汇总", "打卡明细", "人工调整", "暂停记录"]
        summary = workbook["用户汇总"]
        data = dict(zip([c.value for c in summary[1]], [c.value for c in summary[2]], strict=True))
        assert data["QQ号"] == "1001" and data["学号"] == "2026123456"
        assert isinstance(data["绑定时间"], datetime)
        assert data["绑定时间"].hour == 10
        assert data["本周自动计次"] == 1
        assert data["本周人工计次"] == 1
        assert data["本周有效计次"] == 2
        assert data["本周状态"] == "已完成"
        assert data["有效材料数"] == 1
        assert data["累计周有效计次"] == 2
        assert data["完成周数"] == 1
        assert workbook["本周总览"]["B5"].value == 1
        assert workbook["本周总览"]["F5"].value == 1
        assert workbook["本周总览"]._charts
        assert workbook["打卡明细"].max_row == 2
        assert workbook["人工调整"].max_row == 2
        workbook.close()

    asyncio.run(scenario())
