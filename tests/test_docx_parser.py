import zipfile

import pytest
from astrbot_plugin_direct_checkin.services.docx_parser import (
    DocxValidationError,
    evidence_text,
    extract_content,
    validate_docx,
)


def _make_docx(path):
    import docx

    document = docx.Document()
    document.add_heading("学习打卡", level=1)
    document.add_paragraph("今天实现了文件上传接口，并补充了空文件与超限测试。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "用例"
    table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "上传正常"
    table.cell(1, 1).text = "通过"
    document.save(str(path))


def test_validate_and_extract(tmp_path):
    path = tmp_path / "a.docx"
    _make_docx(path)
    validate_docx(path, 10 * 1024 * 1024)
    content = extract_content(path)
    assert "文件上传接口" in content.text
    assert content.table_count == 1
    assert content.char_count > 0
    assert content.headings
    assert "字符数=" in evidence_text(content)


def test_validate_rejects_non_zip(tmp_path):
    path = tmp_path / "bad.docx"
    path.write_text("this is not a docx", encoding="utf-8")
    with pytest.raises(DocxValidationError):
        validate_docx(path, 10 * 1024 * 1024)


def test_validate_rejects_missing_structure(tmp_path):
    path = tmp_path / "empty.docx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("hello.txt", "hi")
    with pytest.raises(DocxValidationError):
        validate_docx(path, 10 * 1024 * 1024)


def test_validate_rejects_zip_bomb(tmp_path):
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<xml/>")
        zf.writestr("word/document.xml", "<xml/>")
        zf.writestr("word/media/big.bin", b"\x00" * 6_000_000)
    with pytest.raises(DocxValidationError):
        validate_docx(path, 10 * 1024 * 1024)
