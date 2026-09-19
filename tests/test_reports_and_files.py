import asyncio
import logging
import warnings
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from astrbot_plugin_direct_checkin.repositories.admin_repo import AdminRepository
from astrbot_plugin_direct_checkin.repositories.db import Database
from astrbot_plugin_direct_checkin.repositories.pause_repo import PauseRepository
from astrbot_plugin_direct_checkin.repositories.submission_repo import SubmissionRepository
from astrbot_plugin_direct_checkin.repositories.user_repo import UserRepository
from astrbot_plugin_direct_checkin.services.ai_review_service import (
    AIReviewService,
    normalize_ai_data,
)
from astrbot_plugin_direct_checkin.services.docx_parser import extract_content, text_for_ai
from astrbot_plugin_direct_checkin.services.export_service import ExportService
from astrbot_plugin_direct_checkin.services.file_service import (
    FileService,
    FileServiceError,
    RepliedFile,
)
from astrbot_plugin_direct_checkin.services.stat_service import StatService
from astrbot_plugin_direct_checkin.tests.test_admin_reset import _make_service
from astrbot_plugin_direct_checkin.utils.time_utils import now_utc, utc_iso, week_key_now
from openpyxl import load_workbook


async def setup(root, target=2):
    db = Database(root / "db")
    await db.initialize()
    repositories = dict(
        user_repo=UserRepository(db),
        submission_repo=SubmissionRepository(db),
        admin_repo=AdminRepository(db),
        pause_repo=PauseRepository(db),
        weekly_limit=target,
        timezone="Asia/Shanghai",
    )
    return (
        db,
        repositories,
        ExportService(**repositories, export_dir=root / "exports"),
        StatService(**repositories, image_dir=root / "images", logger=logging.getLogger("test")),
    )


@pytest.mark.parametrize("target", [1, 2, 3])
def test_chart_export_counts_identity_types_and_unique_paths(tmp_path, target):
    async def run():
        _, repos, exporter, stat = await setup(tmp_path, target)
        now = utc_iso(now_utc())
        week = week_key_now()
        first = await repos["user_repo"].create(
            "000000000000000001", "0020260001", "=1+1", None, now
        )
        second = await repos["user_repo"].create(
            "000000000000000002", "0020260002", "名字比较长的另一位示例同学", None, now
        )
        for _ in range(target):
            await repos["admin_repo"].allocate_positive_adjustment(
                first.id, week, week, target, "9", now
            )
        if target > 1:
            await repos["admin_repo"].allocate_positive_adjustment(
                second.id, week, week, target, "9", now
            )
        snapshot = await stat.report.snapshot()
        assert snapshot.counts == (1, int(target > 1), int(target == 1))
        with warnings.catch_warnings():
            warnings.filterwarnings("error", message="Glyph .* missing")
            first_chart = stat.render(snapshot)
            second_chart = stat.render(snapshot)
        assert first_chart != second_chart and first_chart.read_bytes().startswith(b"\x89PNG")
        all_report, personal = await asyncio.gather(exporter.export(), exporter.export(first.id))
        assert all_report.path != personal.path
        workbook = load_workbook(all_report.path)
        assert workbook["用户汇总"]["A2"].value == "000000000000000001"
        assert workbook["用户汇总"]["B2"].value == "0020260001"
        assert workbook["用户汇总"]["C2"].value == "=1+1"
        assert workbook["用户汇总"]["C2"].data_type == "s"
        assert not any(c.data_type == "f" for sheet in workbook for row in sheet for c in row)
        assert workbook["本周总览"]["B5"].value == snapshot.total
        assert workbook["本周总览"]["D5"].value == snapshot.counts[0]
        assert workbook["本周总览"]["F5"].value == 0.5
        workbook.close()
        workbook = load_workbook(personal.path)
        assert workbook["用户汇总"].max_row == 2
        assert workbook["本周总览"]["B5"].value == 1
        assert workbook["人工调整"].max_row == target + 1
        workbook.close()

    asyncio.run(run())


def test_empty_paused_and_renderer_failure_fallback(tmp_path, monkeypatch):
    async def run():
        _, repos, exporter, stat = await setup(tmp_path)
        empty = await stat.stat()
        assert empty.image_path and empty.image_path.exists()
        assert (await exporter.export()).empty
        now = utc_iso(now_utc())
        week = week_key_now()
        await repos["user_repo"].create("1", "20260001", "A", None, now)
        await repos["pause_repo"].create(
            "week", now, "2099-01-01T00:00:00+00:00", "考试周，暂停原因较长。" * 15, "9", now, week
        )
        paused = await stat.stat()
        assert paused.paused and paused.image_path.exists()
        bundle = await exporter.export()
        workbook = load_workbook(bundle.path)
        assert workbook["本周总览"]["F5"].value == "—"
        assert workbook["用户汇总"]["D2"].value == "暂停周"
        assert workbook["用户汇总"]["F2"].value == 0
        workbook.close()
        monkeypatch.setattr(
            stat, "render", lambda _: (_ for _ in ()).throw(RuntimeError("font unavailable"))
        )
        fallback = await stat.stat()
        assert fallback.image_path is None and "暂停" in fallback.text

    asyncio.run(run())


def test_file_failures_do_not_delete_adapter_source(tmp_path):
    async def run():
        nas = tmp_path / "nas"
        nas.mkdir()
        source = tmp_path / "bad.docx"
        source.write_bytes(b"not a docx")
        component = SimpleNamespace(get_file=AsyncMock(return_value=str(source)))
        service = FileService(str(nas), 1, logging.getLogger("test"))
        user = SimpleNamespace(student_id="20260001")
        with pytest.raises(FileServiceError):
            await service.prepare(
                RepliedFile(component, "bad.docx", "", "q"), user, "unique", "time"
            )
        assert source.exists() and list(nas.iterdir()) == []
        component.size = 2 * 1024 * 1024
        component.get_file.reset_mock()
        with pytest.raises(FileServiceError):
            await service.prepare(
                RepliedFile(component, "bad.docx", "", "q"), user, "unique", "time"
            )
        component.get_file.assert_not_awaited()

    asyncio.run(run())


def test_word_order_and_middle_key_paragraph_survives(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("开头简介")
    for _ in range(10):
        document.add_paragraph("背景资料内容。" * 80)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "测试：修复上传超限错误的唯一关键段落"
    for _ in range(10):
        document.add_paragraph("背景资料内容。" * 80)
    document.add_paragraph("结尾说明")
    path = tmp_path / "long.docx"
    document.save(path)
    content = extract_content(path)
    assert content.text.index("测试:") < content.text.index("结尾说明")
    assert "\n" in content.text
    assert "唯一关键段落" in text_for_ai(content, 1500)


@pytest.mark.parametrize("value", [float("nan"), float("inf"), True, "0.8"])
def test_invalid_ai_confidence_is_technical_error(value):
    with pytest.raises(ValueError):
        normalize_ai_data({"decision": "pass", "confidence": value})


def test_ai_timeout_is_retryable_error():
    async def run():
        context = SimpleNamespace(
            get_current_chat_provider_id=AsyncMock(return_value="test"),
            llm_generate=AsyncMock(side_effect=TimeoutError()),
        )
        service = AIReviewService(context, "rules", 5, 1, logging.getLogger("test"))
        result = await service.review("umo", "text", "", "")
        assert result.error_code == "E_AI_TIMEOUT" and context.llm_generate.await_count == 2

    asyncio.run(run())


def test_reset_cannot_target_system_root(tmp_path):
    async def run():
        _, _, service = await _make_service(tmp_path, Path(tmp_path.anchor))
        with pytest.raises(ValueError):
            service.reset_service._base()

    asyncio.run(run())


def test_reset_unlinks_outside_symlink_without_following(tmp_path):
    async def run():
        nas = tmp_path / "nas"
        nas.mkdir()
        outside = tmp_path / "outside"
        outside.mkdir()
        (outside / "keep").write_bytes(b"keep")
        try:
            (nas / "link").symlink_to(outside, target_is_directory=True)
        except OSError:
            pytest.skip("当前系统未提供创建符号链接权限")
        _, _, service = await _make_service(tmp_path, nas)
        counts, _ = await service.reset_service.execute()
        assert counts is not None
        assert (outside / "keep").read_bytes() == b"keep"
        assert not (nas / "link").exists()

    asyncio.run(run())


def test_all_completed_and_resumed_week_agree_with_excel(tmp_path):
    async def run():
        _, repos, exporter, stat = await setup(tmp_path, 3)
        now = utc_iso(now_utc())
        week = week_key_now()
        user = await repos["user_repo"].create("1", "20260001", "A", None, now)
        for _ in range(3):
            await repos["admin_repo"].allocate_positive_adjustment(user.id, week, week, 3, "9", now)
        snapshot = await stat.report.snapshot()
        assert snapshot.counts == (1, 0, 0)
        assert stat.render(snapshot).is_file()
        await repos["pause_repo"].create(
            "week", now, "2099-01-01T00:00:00+00:00", "暂停", "9", now, week
        )
        assert (await stat.report.snapshot()).paused
        await repos["pause_repo"].resume_active(now)
        resumed = await stat.report.snapshot()
        assert not resumed.paused and resumed.counts == (1, 0, 0)
        assert resumed.users[0].completed_weeks == 1 and resumed.users[0].exempt_weeks == 0
        bundle = await exporter.export()
        workbook = load_workbook(bundle.path)
        assert workbook["本周总览"]["F5"].value == 1
        assert workbook["用户汇总"]["D2"].value == "已完成"
        workbook.close()

    asyncio.run(run())
